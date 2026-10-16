from __future__ import annotations

import base64
import hashlib
import hmac
import io
import mimetypes
import numbers
import re
import unicodedata
from pathlib import Path

import pandas as pd
import plotly.graph_objects as go
import streamlit as st


# ============================================================
# CONFIGURAÇÕES
# ============================================================

BASE_DIR = Path(__file__).resolve().parent

LOGO_DIR = BASE_DIR / "assets" / "logos"
DEFAULT_EXCEL = BASE_DIR / "data" / "dados.xlsx"

DASHBOARD_LOGO_CANDIDATES = [
    BASE_DIR / "assets" / "red_bull_bragantino.png",
    BASE_DIR / "assets" / "redbull_bragantino.png",
    BASE_DIR / "assets" / "red_bull_bragantino.jpg",
    BASE_DIR / "assets" / "redbull_bragantino.jpg",
    BASE_DIR / "assets" / "red_bull_bragantino.jpeg",
    BASE_DIR / "assets" / "redbull_bragantino.jpeg",
    BASE_DIR / "assets" / "red_bull_bragantino.svg",
    BASE_DIR / "assets" / "redbull_bragantino.svg",
    BASE_DIR / "assets" / "logos" / "red_bull_bragantino.png",
    BASE_DIR / "assets" / "logos" / "redbull_bragantino.png",
    BASE_DIR / "assets" / "logos" / "RedBullBragantino.png",
]


# Faixa do gráfico reservada para os logos.
# As barras ocupam somente esta área vertical.
BAR_DOMAIN_BOTTOM = 0.22
BAR_DOMAIN_TOP = 0.82

# Posição vertical dos logos, em coordenadas do papel.
LOGO_Y_POSITION = 0.10
LOGO_SIZE_Y = 0.12


# ============================================================
# PALETA RED BULL / RED BULL BRAGANTINO
# ============================================================

BRAGANTINO_NAVY = "#0B1F4D"
BRAGANTINO_DARK_BLUE = "#243866"
BRAGANTINO_BLUE = "#5579A8"
BRAGANTINO_LIGHT_BLUE = "#8EA6C9"
BRAGANTINO_PALE_BLUE = "#C5CBD5"

BRAGANTINO_RED = "#D71945"
BRAGANTINO_DARK_RED = "#A3143E"

BRAGANTINO_YELLOW = "#FFC629"

BRAGANTINO_BROWN = "#9A6B4F"
BRAGANTINO_GRAY = "#7A8496"
BRAGANTINO_DARK_GRAY = "#4D5A70"
BRAGANTINO_LIGHT_GRAY = "#D1D5DB"


# Nenhuma cor verde é utilizada.
BRAND_PALETTE = [
    BRAGANTINO_NAVY,
    BRAGANTINO_RED,
    BRAGANTINO_YELLOW,
    BRAGANTINO_BLUE,
    BRAGANTINO_LIGHT_BLUE,
    BRAGANTINO_DARK_BLUE,
    BRAGANTINO_DARK_RED,
    BRAGANTINO_BROWN,
    BRAGANTINO_GRAY,
    BRAGANTINO_DARK_GRAY,
    BRAGANTINO_PALE_BLUE,
    BRAGANTINO_LIGHT_GRAY,
]


PROJECTION_COLOR_MAP = {
    "relacionamento": BRAGANTINO_BLUE,
    "socio": BRAGANTINO_NAVY,
    "visitante": BRAGANTINO_YELLOW,
    "ingressos": BRAGANTINO_RED,
}


PLAN_COLOR_MAP = {
    "vip": BRAGANTINO_DARK_BLUE,
    "proximageracao": BRAGANTINO_LIGHT_BLUE,
    "crossconnectionasaprata": BRAGANTINO_BLUE,
    "crossconectionasaprata": BRAGANTINO_BLUE,
    "asasprata": BRAGANTINO_PALE_BLUE,
    "asasplatina": BRAGANTINO_NAVY,
    "asasouro": BRAGANTINO_YELLOW,
    "asasdiamante": BRAGANTINO_RED,
    "asasbronze": BRAGANTINO_BROWN,
    "funcionario": BRAGANTINO_DARK_RED,
    "funcionarios": BRAGANTINO_DARK_RED,
}


PROJECTION_TYPE_ORDER = [
    "Relacionamento",
    "Sócio",
    "Visitante",
    "Ingressos",
]


# ============================================================
# CONFIGURAÇÃO DAS COLUNAS
# ============================================================

PROJECTION_COLUMNS = {
    "ID": [
        "id",
    ],
    "Adversário": [
        "adversario",
        "oponente",
        "time",
    ],
    "Tipo": [
        "tipo",
    ],
    "Quantidade": [
        "quantidade",
        "qtd",
        "qtdade",
        "valor",
    ],
}


ID_COLUMNS = {
    "ID": [
        "id",
    ],
    "Adversário": [
        "adversario",
        "oponente",
        "time",
    ],
}


SOCIOS_COLUMNS = {
    "ID": [
        "id",
    ],
    "Plano": [
        "plano",
    ],
    "Quantidade": [
        "quantidade",
        "qtd",
        "qtdade",
        "valor",
    ],
}


# ============================================================
# AUTENTICAÇÃO
# ============================================================

DEFAULT_AUTH_ITERATIONS = 310_000


def load_auth_config():
    """
    Carrega usuário e hash da senha dos Secrets.

    A senha original não fica armazenada no código.
    """

    try:
        auth_config = st.secrets["auth"]

        configured_username = str(
            auth_config["username"]
        ).strip()

        salt = str(
            auth_config["salt"]
        ).strip()

        password_hash = str(
            auth_config["password_hash"]
        ).strip()

        try:
            iterations = int(
                auth_config["iterations"]
            )
        except Exception:
            iterations = DEFAULT_AUTH_ITERATIONS

        if not configured_username:
            raise ValueError

        if not salt:
            raise ValueError

        if not password_hash:
            raise ValueError

        if iterations < 100_000:
            raise ValueError

        return (
            configured_username,
            salt,
            password_hash,
            iterations,
        )

    except Exception as error:
        st.error(
            "A autenticação não está configurada. "
            "Adicione a seção [auth] nos Secrets "
            "do Streamlit."
        )

        st.stop()


def verify_password(
    password: str,
    salt: str,
    expected_hash: str,
    iterations: int,
) -> bool:
    """
    Calcula o hash da senha digitada e compara
    com o hash salvo nos Secrets.
    """

    calculated_hash = hashlib.pbkdf2_hmac(
        "sha256",
        password.encode("utf-8"),
        salt.encode("utf-8"),
        iterations,
    ).hex()

    return hmac.compare_digest(
        calculated_hash.lower(),
        expected_hash.lower(),
    )


def require_authentication():
    """
    Exibe a tela de login e bloqueia o dashboard
    até o usuário ser autenticado.
    """

    if st.session_state.get(
        "authenticated",
        False,
    ):
        return

    (
        configured_username,
        salt,
        expected_hash,
        iterations,
    ) = load_auth_config()

    st.markdown(
        """
        <div style="
            max-width: 480px;
            margin: 90px auto 25px auto;
            text-align: center;
        ">
            <h1>🔒 Acesso restrito</h1>
            <p>
                Informe suas credenciais para acessar
                o dashboard.
            </p>
        </div>
        """,
        unsafe_allow_html=True,
    )

    with st.form(
        "login_form",
        clear_on_submit=False,
    ):

        username = st.text_input(
            "Usuário",
            autocomplete="username",
        )

        password = st.text_input(
            "Senha",
            type="password",
            autocomplete="current-password",
        )

        submitted = st.form_submit_button(
            "Entrar",
            type="primary",
            width="stretch",
        )

    if submitted:

        username_is_valid = hmac.compare_digest(
            username.strip(),
            configured_username,
        )

        password_is_valid = verify_password(
            password=password,
            salt=salt,
            expected_hash=expected_hash,
            iterations=iterations,
        )

        if (
            username_is_valid
            and password_is_valid
        ):

            st.session_state[
                "authenticated"
            ] = True

            st.rerun()

        else:

            st.error(
                "Usuário ou senha inválidos."
            )

    st.stop()


def render_logout_button():
    """
    Exibe o botão para encerrar a sessão atual.
    """

    if st.sidebar.button(
        "Sair",
        width="stretch",
    ):

        st.session_state[
            "authenticated"
        ] = False

        st.rerun()


# ============================================================
# NORMALIZAÇÃO E TRATAMENTO
# ============================================================

def normalize_key(value) -> str:
    if value is None:
        return ""

    try:
        if pd.isna(value):
            return ""
    except (TypeError, ValueError):
        pass

    text = str(value)

    text = unicodedata.normalize(
        "NFKD",
        text,
    )

    text = "".join(
        character
        for character in text
        if not unicodedata.combining(character)
    )

    text = text.lower()

    text = re.sub(
        r"[^a-z0-9]+",
        "",
        text,
    )

    return text


def normalize_id(value) -> str:
    if value is None:
        return ""

    try:
        if pd.isna(value):
            return ""
    except (TypeError, ValueError):
        pass

    text = str(value).strip()

    if text.endswith(".0"):
        text = text[:-2]

    return text


def clean_text(series: pd.Series) -> pd.Series:
    return (
        series
        .fillna("")
        .astype(str)
        .str.strip()
    )


def round_positive_number(value) -> int:
    """
    Arredonda quantidades positivas.

    735,49 -> 735
    735,50 -> 736
    735,89 -> 736
    """

    return int(float(value) + 0.5)


def parse_number(value):
    try:
        if pd.isna(value):
            return None
    except (TypeError, ValueError):
        pass

    if isinstance(value, numbers.Number):
        return float(value)

    text = str(value).strip()
    text = text.replace(" ", "")

    if text == "":
        return None

    # 1.842,50 -> 1842.50
    if "," in text and "." in text:
        text = text.replace(".", "")
        text = text.replace(",", ".")

    # 1842,50 -> 1842.50
    elif "," in text:
        text = text.replace(",", ".")

    # 1.842 -> 1842
    elif re.fullmatch(
        r"\d{1,3}(?:\.\d{3})+",
        text,
    ):
        text = text.replace(".", "")

    try:
        return float(text)

    except ValueError:
        return None


def parse_quantity(series: pd.Series) -> pd.Series:
    values = series.map(parse_number)

    values = pd.to_numeric(
        values,
        errors="coerce",
    )

    if values.isna().any():

        invalid_indexes = list(
            values[values.isna()].index[:5]
        )

        raise ValueError(
            "A coluna Quantidade possui valores "
            "vazios ou inválidos. "
            f"Exemplos de linhas: {invalid_indexes}"
        )

    return (
        values
        .map(round_positive_number)
        .astype("int64")
    )


# ============================================================
# PREPARAÇÃO DAS TABELAS
# ============================================================

def canonicalize_columns(
    raw_df: pd.DataFrame,
    required_columns: dict,
    source_name: str,
) -> pd.DataFrame:

    if raw_df is None or raw_df.empty:
        raise ValueError(
            f"A fonte '{source_name}' está vazia."
        )

    df = raw_df.dropna(
        how="all"
    ).copy()

    normalized_columns = {
        column: normalize_key(column)
        for column in df.columns
    }

    rename_map = {}

    for canonical_name, aliases in required_columns.items():

        normalized_aliases = {
            normalize_key(alias)
            for alias in aliases
        }

        found_column = None

        for (
            original_column,
            normalized_column,
        ) in normalized_columns.items():

            if normalized_column in normalized_aliases:
                found_column = original_column
                break

        if found_column is None:

            columns_found = ", ".join(
                str(column)
                for column in df.columns
            )

            raise ValueError(
                f"A fonte '{source_name}' não possui "
                f"a coluna obrigatória "
                f"'{canonical_name}'.\n\n"
                f"Colunas encontradas: {columns_found}"
            )

        rename_map[found_column] = canonical_name

    df = df.rename(
        columns=rename_map
    )

    return df[
        list(required_columns.keys())
    ]


def validate_nonempty(
    df: pd.DataFrame,
    column: str,
    source_name: str,
):
    invalid_rows = df[column].eq("")

    if invalid_rows.any():

        raise ValueError(
            f"A fonte '{source_name}' possui "
            f"valores vazios na coluna '{column}'."
        )


def prepare_projection(
    raw_df: pd.DataFrame,
    source_name: str,
) -> pd.DataFrame:

    df = canonicalize_columns(
        raw_df=raw_df,
        required_columns=PROJECTION_COLUMNS,
        source_name=source_name,
    )

    df["ID"] = df["ID"].map(
        normalize_id
    )

    df["Adversário"] = clean_text(
        df["Adversário"]
    )

    df["Tipo"] = clean_text(
        df["Tipo"]
    )

    df["Quantidade"] = parse_quantity(
        df["Quantidade"]
    )

    validate_nonempty(
        df,
        "ID",
        source_name,
    )

    validate_nonempty(
        df,
        "Adversário",
        source_name,
    )

    validate_nonempty(
        df,
        "Tipo",
        source_name,
    )

    return df


def prepare_ids(
    raw_df: pd.DataFrame,
    source_name: str,
) -> pd.DataFrame:

    df = canonicalize_columns(
        raw_df=raw_df,
        required_columns=ID_COLUMNS,
        source_name=source_name,
    )

    df["ID"] = df["ID"].map(
        normalize_id
    )

    df["Adversário"] = clean_text(
        df["Adversário"]
    )

    validate_nonempty(
        df,
        "ID",
        source_name,
    )

    validate_nonempty(
        df,
        "Adversário",
        source_name,
    )

    return df


def prepare_socios(
    raw_df: pd.DataFrame,
    source_name: str,
) -> pd.DataFrame:

    df = canonicalize_columns(
        raw_df=raw_df,
        required_columns=SOCIOS_COLUMNS,
        source_name=source_name,
    )

    df["ID"] = df["ID"].map(
        normalize_id
    )

    df["Plano"] = clean_text(
        df["Plano"]
    )

    df["Quantidade"] = parse_quantity(
        df["Quantidade"]
    )

    validate_nonempty(
        df,
        "ID",
        source_name,
    )

    validate_nonempty(
        df,
        "Plano",
        source_name,
    )

    return df


# ============================================================
# LEITURA DO EXCEL
# ============================================================

@st.cache_data(
    show_spinner=False,
)
def read_excel_sources(
    file_bytes: bytes,
) -> dict[str, pd.DataFrame]:

    sources = {}

    sheets = pd.read_excel(
        io.BytesIO(file_bytes),
        sheet_name=None,
        engine="openpyxl",
    )

    # Lê tabelas estruturadas criadas com Ctrl + T.
    try:
        from openpyxl import load_workbook

        workbook = load_workbook(
            io.BytesIO(file_bytes),
            read_only=False,
            data_only=True,
        )

        for worksheet in workbook.worksheets:

            for table_name in worksheet.tables:

                table = worksheet.tables[
                    table_name
                ]

                rows = worksheet[
                    table.ref
                ]

                values = [
                    [
                        cell.value
                        for cell in row
                    ]
                    for row in rows
                ]

                if len(values) < 2:
                    continue

                headers = []

                for index, header in enumerate(
                    values[0]
                ):

                    if header is None:

                        headers.append(
                            f"coluna_{index}"
                        )

                    else:

                        headers.append(
                            str(header).strip()
                        )

                table_df = pd.DataFrame(
                    values[1:],
                    columns=headers,
                )

                source_label = (
                    f"Tabela: {table_name}"
                )

                if source_label in sources:

                    source_label = (
                        f"Tabela: {table_name} "
                        f"({worksheet.title})"
                    )

                sources[source_label] = table_df

    except Exception:
        pass

    # Adiciona as abas.
    for sheet_name, dataframe in sheets.items():

        sources[
            f"Aba: {sheet_name}"
        ] = dataframe

    return sources


# ============================================================
# LEITURA OPCIONAL DO POSTGRESQL
# ============================================================

def get_database_config():
    try:
        return st.secrets["database"]
    except Exception:
        return None


def has_database_config() -> bool:
    config = get_database_config()

    if config is None:
        return False

    if "url" in config:
        return bool(
            str(config["url"]).strip()
        )

    required_fields = [
        "host",
        "name",
        "user",
        "password",
    ]

    return all(
        field in config
        and str(config[field]).strip()
        for field in required_fields
    )


@st.cache_resource(
    show_spinner=False,
)
def get_database_engine():
    """
    Cria uma conexão com PostgreSQL usando
    credenciais em st.secrets.
    """

    try:
        from sqlalchemy import URL, create_engine

    except ImportError as error:
        raise RuntimeError(
            "SQLAlchemy não está instalado. "
            "Adicione SQLAlchemy e psycopg2-binary "
            "ao requirements.txt."
        ) from error

    config = get_database_config()

    if config is None:
        raise RuntimeError(
            "A seção [database] não foi encontrada "
            "nos Secrets."
        )

    if "url" in config:

        database_url = str(
            config["url"]
        ).strip()

        return create_engine(
            database_url,
            pool_pre_ping=True,
            pool_recycle=1800,
        )

    host = str(
        config["host"]
    ).strip()

    port = int(
        config.get(
            "port",
            5432,
        )
    )

    database_name = str(
        config["name"]
    ).strip()

    username = str(
        config["user"]
    ).strip()

    password = str(
        config["password"]
    )

    database_url = URL.create(
        drivername="postgresql+psycopg2",
        username=username,
        password=password,
        host=host,
        port=port,
        database=database_name,
        query={
            "sslmode": "require",
        },
    )

    return create_engine(
        database_url,
        pool_pre_ping=True,
        pool_recycle=1800,
    )


@st.cache_data(
    ttl=900,
    show_spinner=False,
)
def read_database_sources():
    """
    Lê as tabelas do PostgreSQL.

    Tabelas esperadas:

        public.projecao
        public.id_adversario
        public.socios
    """

    from sqlalchemy import text

    engine = get_database_engine()

    queries = {
        "Banco: Projeção": """
            select
                id,
                adversario,
                tipo,
                quantidade
            from public.projecao
            order by id, tipo
        """,

        "Banco: Id": """
            select
                id,
                adversario
            from public.id_adversario
            order by id
        """,

        "Banco: Sócios": """
            select
                id,
                plano,
                quantidade
            from public.socios
            order by id, plano
        """,
    }

    sources = {}

    with engine.connect() as connection:

        for source_name, query in queries.items():

            dataframe = pd.read_sql_query(
                text(query),
                connection,
            )

            sources[source_name] = dataframe

    return sources


# ============================================================
# SELEÇÃO AUTOMÁTICA DAS FONTES
# ============================================================

def default_source_index(
    source_names: list[str],
    aliases: list[str],
    fallback: int,
) -> int:

    normalized_aliases = {
        normalize_key(alias)
        for alias in aliases
    }

    for index, source_name in enumerate(
        source_names
    ):

        name = source_name.split(
            ": ",
            1,
        )[-1]

        name = name.split(
            " (",
            1,
        )[0]

        normalized_name = normalize_key(
            name
        )

        if (
            normalized_name in normalized_aliases
            or any(
                normalized_name.startswith(alias)
                for alias in normalized_aliases
            )
        ):
            return index

    return min(
        fallback,
        len(source_names) - 1,
    )


# ============================================================
# FORMATAÇÃO
# ============================================================

def format_quantity(value) -> str:
    integer_value = round_positive_number(
        value
    )

    return f"{integer_value:,}".replace(
        ",",
        ".",
    )


def ordered_unique(
    values: list,
    preferred_order: list[str] | None = None,
) -> list:

    unique_values = list(
        dict.fromkeys(values)
    )

    if not preferred_order:
        return unique_values

    result = []

    for preferred_value in preferred_order:

        preferred_key = normalize_key(
            preferred_value
        )

        for value in unique_values:

            if (
                value not in result
                and normalize_key(value)
                == preferred_key
            ):
                result.append(value)

    for value in unique_values:

        if value not in result:
            result.append(value)

    return result


def is_employee_plan(value) -> bool:
    plan_key = normalize_key(value)

    return (
        plan_key in {
            "funcionario",
            "funcionarios",
        }
        or "funcion" in plan_key
    )


# ============================================================
# TEMA
# ============================================================

def get_configured_theme_mode() -> str:
    try:
        configured_theme = st.get_option(
            "theme.base"
        )
    except Exception:
        configured_theme = "light"

    configured_theme = str(
        configured_theme
    ).lower()

    if configured_theme == "dark":
        return "dark"

    return "light"


def get_theme_colors(
    theme_mode: str,
) -> dict[str, str]:

    if theme_mode == "dark":

        return {
            "title": "#FFFFFF",
            "legend": "#FFFFFF",
            "total_label": "#FFFFFF",
        }

    return {
        "title": BRAGANTINO_NAVY,
        "legend": BRAGANTINO_NAVY,
        "total_label": BRAGANTINO_NAVY,
    }


# ============================================================
# CORES DOS SEGMENTOS
# ============================================================

def create_fallback_gray(
    gray_index: int,
) -> str:

    gray_value = 45 + (
        gray_index * 18
    )

    gray_value = min(
        gray_value,
        235,
    )

    return (
        f"#{gray_value:02X}"
        f"{gray_value:02X}"
        f"{gray_value:02X}"
    )


def get_unique_segment_colors(
    segments: list,
    preferred_colors: dict[str, str],
) -> dict:

    """
    Garante que nenhuma cor se repita
    dentro do mesmo gráfico.
    """

    colors_by_segment = {}
    used_colors = set()

    for segment in segments:

        segment_key = normalize_key(
            segment
        )

        preferred_color = preferred_colors.get(
            segment_key
        )

        if (
            preferred_color
            and preferred_color not in used_colors
        ):

            selected_color = preferred_color

        else:

            selected_color = None

            for candidate_color in BRAND_PALETTE:

                if candidate_color not in used_colors:

                    selected_color = candidate_color
                    break

            if selected_color is None:

                gray_index = 0

                while selected_color is None:

                    gray_candidate = create_fallback_gray(
                        gray_index
                    )

                    if gray_candidate not in used_colors:

                        selected_color = (
                            gray_candidate
                        )

                    gray_index += 1

        colors_by_segment[segment] = selected_color
        used_colors.add(selected_color)

    return colors_by_segment


def get_text_color_for_segment(
    background_color: str,
) -> str:

    color = background_color.lstrip("#")

    red = int(
        color[0:2],
        16,
    )

    green = int(
        color[2:4],
        16,
    )

    blue = int(
        color[4:6],
        16,
    )

    luminance = (
        0.299 * red
        + 0.587 * green
        + 0.114 * blue
    )

    if luminance > 170:
        return BRAGANTINO_NAVY

    return "#FFFFFF"


# ============================================================
# LOGOS
# ============================================================

def find_logo(
    adversary: str,
) -> Path | None:

    if not LOGO_DIR.exists():
        return None

    valid_extensions = {
        ".png",
        ".jpg",
        ".jpeg",
        ".svg",
    }

    adversary_key = normalize_key(
        adversary
    )

    for file_path in sorted(
        LOGO_DIR.iterdir()
    ):

        if not file_path.is_file():
            continue

        if (
            file_path.suffix.lower()
            not in valid_extensions
        ):
            continue

        file_key = normalize_key(
            file_path.stem
        )

        if file_key == adversary_key:
            return file_path

    return None


def find_dashboard_logo() -> Path | None:

    for candidate in DASHBOARD_LOGO_CANDIDATES:

        if candidate.exists():
            return candidate

    possible_names = {
        "redbullbragantino",
        "redbullbragantinoescudo",
        "logoredbullbragantino",
        "escudoredbullbragantino",
    }

    possible_directories = [
        BASE_DIR / "assets",
        BASE_DIR / "assets" / "logos",
    ]

    valid_extensions = {
        ".png",
        ".jpg",
        ".jpeg",
        ".svg",
    }

    for directory in possible_directories:

        if not directory.exists():
            continue

        for file_path in directory.iterdir():

            if not file_path.is_file():
                continue

            if (
                file_path.suffix.lower()
                not in valid_extensions
            ):
                continue

            if normalize_key(
                file_path.stem
            ) in possible_names:

                return file_path

    return None


def image_to_data_uri(
    image_path: Path,
) -> str:

    mime_type = (
        mimetypes.guess_type(
            image_path.name
        )[0]
        or "image/png"
    )

    image_bytes = image_path.read_bytes()

    encoded_image = base64.b64encode(
        image_bytes
    ).decode("utf-8")

    return (
        f"data:{mime_type};base64,"
        f"{encoded_image}"
    )


# ============================================================
# GRÁFICO
# ============================================================

def create_stacked_chart(
    dataframe: pd.DataFrame,
    category_column: str,
    segment_column: str,
    value_column: str,
    chart_title: str,
    segment_order: list[str] | None = None,
    segment_color_map: dict[str, str] | None = None,
    chart_height: int = 570,
    show_legend: bool = True,
    theme_mode: str = "light",
) -> go.Figure | None:

    if dataframe.empty:
        return None

    categories = list(
        dict.fromkeys(
            dataframe[
                category_column
            ].tolist()
        )
    )

    segments = ordered_unique(
        dataframe[
            segment_column
        ].tolist(),
        preferred_order=segment_order,
    )

    grouped = (
        dataframe
        .groupby(
            [
                category_column,
                segment_column,
            ],
            sort=False,
            as_index=False,
        )[value_column]
        .sum()
    )

    grouped[value_column] = (
        grouped[value_column]
        .map(round_positive_number)
        .astype("int64")
    )

    pivot = grouped.pivot(
        index=category_column,
        columns=segment_column,
        values=value_column,
    )

    pivot = (
        pivot
        .reindex(
            index=categories,
            columns=segments,
        )
        .fillna(0)
        .astype("int64")
    )

    x_positions = list(
        range(len(categories))
    )

    totals = (
        pivot
        .sum(axis=1)
        .map(round_positive_number)
        .astype("int64")
        .tolist()
    )

    maximum_total = max(
        max(totals),
        1,
    )

    if segment_color_map is None:
        segment_color_map = {}

    segment_colors = get_unique_segment_colors(
        segments=segments,
        preferred_colors=segment_color_map,
    )

    theme_colors = get_theme_colors(
        theme_mode
    )

    figure = go.Figure()

    # ========================================================
    # COLUNAS EMPILHADAS
    # ========================================================

    for segment in segments:

        values = (
            pivot[segment]
            .map(round_positive_number)
            .astype("int64")
            .tolist()
        )

        segment_color = segment_colors[
            segment
        ]

        inside_label_color = (
            get_text_color_for_segment(
                segment_color
            )
        )

        texts = [
            (
                format_quantity(value)
                if value > 0
                else ""
            )
            for value in values
        ]

        custom_data = [
            [category]
            for category in categories
        ]

        hover_template = (
            f"{category_column}: "
            "%{customdata[0]}<br>"
            f"{segment_column}: {segment}<br>"
            "Quantidade: %{y:,.0f}"
            "<extra></extra>"
        )

        figure.add_trace(
            go.Bar(
                x=x_positions,
                y=values,
                name=str(segment),
                text=texts,
                textposition="inside",
                insidetextanchor="middle",
                textfont=dict(
                    size=10,
                    color=inside_label_color,
                    family="Arial, sans-serif",
                ),
                marker=dict(
                    color=segment_color,
                    line=dict(
                        width=0,
                    ),
                ),
                customdata=custom_data,
                hovertemplate=hover_template,
                showlegend=show_legend,
            )
        )

    # ========================================================
    # TOTAIS NO TOPO
    # ========================================================

    for position, total in zip(
        x_positions,
        totals,
    ):

        figure.add_annotation(
            x=position,
            y=total,
            xref="x",
            yref="y",
            text=(
                f"<b>{format_quantity(total)}</b>"
            ),
            showarrow=False,
            yshift=9,
            font=dict(
                size=14,
                color=theme_colors["total_label"],
                family="Arial Black, Arial, sans-serif",
            ),
            align="center",
        )

    # ========================================================
    # LOGOS DOS ADVERSÁRIOS
    # ========================================================
    #
    # Esta faixa está separada da área das barras.
    # As barras começam em BAR_DOMAIN_BOTTOM.
    # Os logos ficam abaixo, em y=0.10.
    #

    for position, category in enumerate(
        categories
    ):

        logo_path = find_logo(
            category
        )

        if logo_path is None:
            continue

        logo_size_x = (
            0.30
            if len(categories) <= 6
            else 0.22
        )

        figure.add_layout_image(
            dict(
                source=image_to_data_uri(
                    logo_path
                ),

                # Alinhamento horizontal com
                # a coluna correspondente.
                x=position,
                xref="x",

                # Faixa inferior exclusiva dos logos.
                y=LOGO_Y_POSITION,
                yref="paper",

                sizex=logo_size_x,
                sizey=LOGO_SIZE_Y,

                xanchor="center",
                yanchor="middle",

                sizing="contain",
                layer="above",
            )
        )

    # ========================================================
    # LAYOUT
    # ========================================================

    figure.update_layout(
        title=dict(
            text=f"<b>{chart_title}</b>",
            x=0,
            xanchor="left",
            y=0.99,
            yanchor="top",
            font=dict(
                size=19,
                color=theme_colors["title"],
                family="Arial, sans-serif",
            ),
        ),

        template="plotly_white",

        # Transparência para funcionar no claro
        # e no escuro.
        paper_bgcolor="rgba(0,0,0,0)",
        plot_bgcolor="rgba(0,0,0,0)",

        font=dict(
            color=theme_colors["legend"],
            family="Arial, sans-serif",
        ),

        barmode="stack",
        height=chart_height,

        # Espaço inferior para a faixa dos logos.
        margin=dict(
            l=20,
            r=30,
            t=150 if show_legend else 105,
            b=90,
        ),

        bargap=0.25,
        showlegend=show_legend,

        # A legenda fica na parte superior,
        # separada da área ocupada pelas barras.
        legend=dict(
            orientation="h",
            y=0.92,
            yanchor="top",
            x=0,
            xanchor="left",
            font=dict(
                size=11,
                color=theme_colors["legend"],
                family="Arial, sans-serif",
            ),
            bgcolor="rgba(0,0,0,0)",
            borderwidth=0,
            traceorder="normal",
        ),

        uniformtext_minsize=8,
        uniformtext_mode="hide",
    )

    # ========================================================
    # EIXO X
    # ========================================================

    # O eixo permanece oculto, mas é usado para
    # alinhar os logos horizontalmente.
    figure.update_xaxes(
        visible=False,
        showticklabels=False,
        showgrid=False,
        zeroline=False,
        showline=False,
        ticks="",
        domain=[
            0,
            1,
        ],
        range=[
            -0.6,
            max(
                len(categories) - 0.4,
                0.6,
            ),
        ],
    )

    # ========================================================
    # EIXO Y
    # ========================================================

    # As barras ocupam somente a faixa entre
    # 22% e 82% da altura do gráfico.
    #
    # A faixa abaixo de 22% fica reservada
    # exclusivamente para os escudos.
    figure.update_yaxes(
        visible=False,
        showticklabels=False,
        showgrid=False,
        zeroline=False,
        showline=False,
        ticks="",
        title_text=None,
        domain=[
            BAR_DOMAIN_BOTTOM,
            BAR_DOMAIN_TOP,
        ],
        rangemode="tozero",
        range=[
            0,
            maximum_total * 1.28,
        ],
    )

    return figure


# ============================================================
# APLICAÇÃO STREAMLIT
# ============================================================

def main():

    st.set_page_config(
        page_title="Dashboard de Adversários",
        page_icon="📊",
        layout="wide",
    )

    # Login antes de carregar os dados.
    require_authentication()

    # --------------------------------------------------------
    # BARRA LATERAL
    # --------------------------------------------------------

    with st.sidebar:

        render_logout_button()

        st.divider()

        st.header(
            "Configurações"
        )

        label_theme_option = st.selectbox(
            "Tema dos rótulos dos gráficos",
            options=[
                "Automático",
                "Claro",
                "Escuro",
            ],
            index=0,
            help=(
                "No modo Escuro, os totais ficam "
                "brancos e em negrito."
            ),
        )

        if label_theme_option == "Claro":

            theme_mode = "light"

        elif label_theme_option == "Escuro":

            theme_mode = "dark"

        else:

            theme_mode = get_configured_theme_mode()

        st.divider()

        uploaded_file = st.file_uploader(
            "Enviar Excel manualmente",
            type=[
                "xlsx",
                "xlsm",
            ],
            help=(
                "Opcional. Se não for enviado um arquivo, "
                "o aplicativo usará o banco configurado "
                "ou o arquivo data/dados.xlsx."
            ),
        )

        database_available = (
            has_database_config()
        )

        if database_available:

            st.info(
                "Fonte principal: banco de dados"
            )

            if st.button(
                "Atualizar dados",
                width="stretch",
            ):

                read_database_sources.clear()
                get_database_engine.clear()

                st.rerun()

        else:

            st.info(
                "Fonte principal: Excel"
            )

    # --------------------------------------------------------
    # CARREGAMENTO DOS DADOS
    # --------------------------------------------------------

    try:

        if uploaded_file is not None:

            excel_bytes = (
                uploaded_file.getvalue()
            )

            sources = read_excel_sources(
                excel_bytes
            )

            data_source_label = (
                f"Excel enviado: "
                f"{uploaded_file.name}"
            )

        elif database_available:

            with st.spinner(
                "Carregando dados do banco..."
            ):

                sources = read_database_sources()

            data_source_label = (
                "Banco de dados"
            )

        elif DEFAULT_EXCEL.exists():

            excel_bytes = (
                DEFAULT_EXCEL.read_bytes()
            )

            sources = read_excel_sources(
                excel_bytes
            )

            data_source_label = (
                f"Arquivo padrão: "
                f"{DEFAULT_EXCEL.name}"
            )

        else:

            sources = {}
            data_source_label = ""

    except Exception as error:

        st.error(
            "Não foi possível carregar os dados."
        )

        st.exception(error)
        st.stop()

    if not sources:

        st.warning(
            "Nenhuma fonte de dados foi encontrada. "
            "Envie um Excel, configure o banco ou "
            "coloque data/dados.xlsx."
        )

        st.stop()

    # --------------------------------------------------------
    # CABEÇALHO
    # --------------------------------------------------------

    header_title, header_logo = st.columns(
        [9, 1]
    )

    with header_title:

        st.title(
            "📊 Dashboard de projeção e sócios"
        )

        st.caption(
            "Projeção e acompanhamento de público "
            "para os próximos jogos"
        )

        if data_source_label:

            st.caption(
                f"Fonte: {data_source_label}"
            )

    with header_logo:

        dashboard_logo = find_dashboard_logo()

        if dashboard_logo is not None:

            st.image(
                str(dashboard_logo),
                width=110,
            )

    # --------------------------------------------------------
    # SELEÇÃO DAS FONTES
    # --------------------------------------------------------

    source_names = list(
        sources.keys()
    )

    with st.sidebar:

        st.divider()

        st.subheader(
            "Fontes das tabelas"
        )

        projection_source = st.selectbox(
            "Tabela ou aba da Projeção",
            source_names,
            index=default_source_index(
                source_names,
                [
                    "projecao",
                    "projecoes",
                    "projection",
                ],
                0,
            ),
        )

        id_source = st.selectbox(
            "Tabela ou aba de ID",
            source_names,
            index=default_source_index(
                source_names,
                [
                    "id",
                    "ids",
                    "identificacao",
                ],
                min(
                    1,
                    len(source_names) - 1,
                ),
            ),
        )

        socios_source = st.selectbox(
            "Tabela ou aba de Sócios",
            source_names,
            index=default_source_index(
                source_names,
                [
                    "socios",
                    "socio",
                ],
                min(
                    2,
                    len(source_names) - 1,
                ),
            ),
        )

    # --------------------------------------------------------
    # PREPARAÇÃO DAS TABELAS
    # --------------------------------------------------------

    try:

        projection = prepare_projection(
            sources[projection_source],
            projection_source,
        )

        id_table = prepare_ids(
            sources[id_source],
            id_source,
        )

        socios = prepare_socios(
            sources[socios_source],
            socios_source,
        )

    except Exception as error:

        st.error(
            "Não foi possível preparar as tabelas."
        )

        st.exception(error)
        st.stop()

    # --------------------------------------------------------
    # RELACIONAMENTO ID X ADVERSÁRIO
    # --------------------------------------------------------

    conflicting_ids = (
        id_table
        .groupby("ID")[
            "Adversário"
        ]
        .nunique()
    )

    conflicting_ids = conflicting_ids[
        conflicting_ids > 1
    ]

    if not conflicting_ids.empty:

        st.warning(
            "Existem IDs associados a mais "
            "de um adversário. Será utilizado "
            "o primeiro adversário encontrado "
            "para cada ID."
        )

    id_lookup = (
        id_table
        .drop_duplicates("ID")
        .set_index("ID")[
            "Adversário"
        ]
    )

    socios["Adversário"] = (
        socios["ID"].map(id_lookup)
    )

    unknown_ids_mask = (
        socios["Adversário"].isna()
        | socios["Adversário"].eq("")
    )

    if unknown_ids_mask.any():

        unknown_id_list = (
            socios.loc[
                unknown_ids_mask,
                "ID",
            ]
            .drop_duplicates()
            .tolist()
        )

        st.warning(
            "Os seguintes IDs da tabela Sócios "
            "não foram encontrados na tabela Id "
            f"e serão ignorados: {unknown_id_list}"
        )

        socios = socios.loc[
            ~unknown_ids_mask
        ].copy()

    # --------------------------------------------------------
    # FILTRO DE ADVERSÁRIOS
    # --------------------------------------------------------

    all_adversaries = list(
        dict.fromkeys(
            projection[
                "Adversário"
            ].tolist()
            + socios[
                "Adversário"
            ].tolist()
        )
    )

    with st.sidebar:

        st.divider()

        selected_adversaries = st.multiselect(
            "Filtrar adversários",
            options=all_adversaries,
            default=all_adversaries,
        )

    if not selected_adversaries:

        st.warning(
            "Selecione pelo menos um adversário."
        )

        st.stop()

    projection_filtered = projection[
        projection[
            "Adversário"
        ].isin(selected_adversaries)
    ].copy()

    socios_filtered = socios[
        socios[
            "Adversário"
        ].isin(selected_adversaries)
    ].copy()

    # --------------------------------------------------------
    # SEPARAÇÃO DOS FUNCIONÁRIOS
    # --------------------------------------------------------

    employee_mask = socios_filtered[
        "Plano"
    ].map(
        is_employee_plan
    )

    socios_employee = socios_filtered[
        employee_mask
    ].copy()

    socios_regular = socios_filtered[
        ~employee_mask
    ].copy()

    # --------------------------------------------------------
    # KPIs
    # --------------------------------------------------------

    total_projection = (
        projection_filtered[
            "Quantidade"
        ].sum()
    )

    total_socios = (
        socios_filtered[
            "Quantidade"
        ].sum()
    )

    number_of_games = (
        projection_filtered[
            "Adversário"
        ].nunique()
    )

    if number_of_games > 0:

        average_projection = (
            total_projection
            / number_of_games
        )

    else:

        average_projection = 0

    number_of_adversaries = len(
        selected_adversaries
    )

    number_of_plans = (
        socios_filtered[
            "Plano"
        ].nunique()
    )

    col1, col2, col3, col4, col5 = st.columns(
        5
    )

    col1.metric(
        "Total projetado",
        format_quantity(
            total_projection
        ),
    )

    col2.metric(
        "Total de sócios",
        format_quantity(
            total_socios
        ),
    )

    col3.metric(
        f"Média dos {number_of_games} jogos",
        format_quantity(
            average_projection
        ),
    )

    col4.metric(
        "Adversários",
        number_of_adversaries,
    )

    col5.metric(
        "Planos",
        number_of_plans,
    )

    st.divider()

    # ========================================================
    # GRÁFICO DE PROJEÇÃO
    # ========================================================

    st.subheader(
        "Projeção por adversário"
    )

    st.caption(
        "Colunas empilhadas por tipo. "
        "O número em negrito no topo representa "
        "o total."
    )

    projection_chart = create_stacked_chart(
        dataframe=projection_filtered,
        category_column="Adversário",
        segment_column="Tipo",
        value_column="Quantidade",
        chart_title="Projeção total por adversário",
        segment_order=PROJECTION_TYPE_ORDER,
        segment_color_map=PROJECTION_COLOR_MAP,
        chart_height=570,
        show_legend=True,
        theme_mode=theme_mode,
    )

    if projection_chart is not None:

        st.plotly_chart(
            projection_chart,
            width="stretch",
            config={
                "displaylogo": False,
            },
        )

    else:

        st.info(
            "Não há dados de projeção."
        )

    # ========================================================
    # GRÁFICO DE SÓCIOS
    # ========================================================

    st.subheader(
        "Sócios por adversário"
    )

    st.caption(
        "Colunas empilhadas por plano. "
        "Funcionários aparecem separadamente."
    )

    socios_chart = create_stacked_chart(
        dataframe=socios_regular,
        category_column="Adversário",
        segment_column="Plano",
        value_column="Quantidade",
        chart_title="Sócios por adversário e plano",
        segment_color_map=PLAN_COLOR_MAP,
        chart_height=620,
        show_legend=True,
        theme_mode=theme_mode,
    )

    if socios_chart is not None:

        st.plotly_chart(
            socios_chart,
            width="stretch",
            config={
                "displaylogo": False,
            },
        )

    else:

        st.info(
            "Não há dados de sócios regulares."
        )

    # ========================================================
    # GRÁFICO DE FUNCIONÁRIOS
    # ========================================================

    st.subheader(
        "Plano Funcionários por partida"
    )

    st.caption(
        "Quantidade de funcionários por adversário."
    )

    employee_chart = create_stacked_chart(
        dataframe=socios_employee,
        category_column="Adversário",
        segment_column="Plano",
        value_column="Quantidade",
        chart_title="Funcionários por adversário",
        segment_order=[
            "Funcionários",
        ],
        segment_color_map=PLAN_COLOR_MAP,
        chart_height=400,
        show_legend=False,
        theme_mode=theme_mode,
    )

    if employee_chart is not None:

        st.plotly_chart(
            employee_chart,
            width="stretch",
            config={
                "displaylogo": False,
            },
        )

    else:

        st.info(
            "Não há registros de funcionários."
        )

    # ========================================================
    # AVISOS SOBRE LOGOS
    # ========================================================

    logos_missing = [
        adversary
        for adversary in selected_adversaries
        if find_logo(adversary) is None
    ]

    if logos_missing:

        st.info(
            "Logos não encontrados em "
            "assets/logos para: "
            + ", ".join(logos_missing)
        )

    if find_dashboard_logo() is None:

        st.info(
            "Logo do Red Bull Bragantino não encontrado. "
            "Salve o arquivo como "
            "'assets/red_bull_bragantino.png'."
        )

    # ========================================================
    # DADOS TRATADOS
    # ========================================================

    with st.expander(
        "Visualizar dados tratados"
    ):

        st.write(
            "Projeção"
        )

        st.dataframe(
            projection_filtered,
            width="stretch",
            hide_index=True,
        )

        st.write(
            "Sócios regulares"
        )

        st.dataframe(
            socios_regular,
            width="stretch",
            hide_index=True,
        )

        st.write(
            "Funcionários"
        )

        st.dataframe(
            socios_employee,
            width="stretch",
            hide_index=True,
        )

        st.write(
            "Relacionamento ID x adversário"
        )

        st.dataframe(
            id_table,
            width="stretch",
            hide_index=True,
        )


if __name__ == "__main__":
    main()
